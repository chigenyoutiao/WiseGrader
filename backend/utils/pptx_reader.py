import zipfile
from pathlib import Path
from typing import Dict, Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


P_MEDIA_EXTS = {'.mp4', '.mov', '.avi', '.wmv', '.mkv', '.mp3', '.wav', '.aac', '.m4a'}
NSMAP = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


def _count_media_files(pptx_path: Path) -> int:
    try:
        with zipfile.ZipFile(pptx_path, "r") as zf:
            return sum(
                1 for name in zf.namelist()
                if name.startswith("ppt/media/") and Path(name).suffix.lower() in P_MEDIA_EXTS
            )
    except Exception:
        return 0


def read_pptx(file_path) -> Dict[str, Any]:
    """
    提取 PPTX 文本与客观指标，用于评分参考。
    返回: {"text": str, "metrics": {...}}
    """
    pptx_path = Path(file_path)
    prs = Presentation(pptx_path)

    text_parts = []
    metrics = {
        "slide_count": len(prs.slides),
        "text_shape_count": 0,
        "picture_count": 0,
        "table_count": 0,
        "chart_count": 0,
        "media_count": _count_media_files(pptx_path),
        "blank_slide_count": 0,
        "transition_slides": 0,
        "timing_slides": 0,
        "has_animation_or_transition": False,
    }

    for slide in prs.slides:
        slide_has_content = False

        # 动画/切换信号（粗粒度）
        try:
            if slide._element.xpath(".//p:transition", namespaces=NSMAP):
                metrics["transition_slides"] += 1
            if slide._element.xpath(".//p:timing", namespaces=NSMAP):
                metrics["timing_slides"] += 1
        except Exception:
            pass

        for shape in slide.shapes:
            stype = shape.shape_type

            # 文本
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                paragraphs = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
                if paragraphs:
                    text_parts.append("\n".join(paragraphs))
                    metrics["text_shape_count"] += 1
                    slide_has_content = True

            # 图片
            if stype == MSO_SHAPE_TYPE.PICTURE:
                metrics["picture_count"] += 1
                slide_has_content = True

            # 表格
            if getattr(shape, "has_table", False) and shape.has_table:
                metrics["table_count"] += 1
                slide_has_content = True

            # 图表
            if getattr(shape, "has_chart", False):
                metrics["chart_count"] += 1
                slide_has_content = True

        if not slide_has_content:
            metrics["blank_slide_count"] += 1

    metrics["has_animation_or_transition"] = (
        metrics["transition_slides"] > 0 or metrics["timing_slides"] > 0
    )

    text_content = "\n\n".join(text_parts).strip()

    return {
        "text": text_content,
        "metrics": metrics,
    }
