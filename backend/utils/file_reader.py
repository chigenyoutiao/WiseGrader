import os
import shutil
import subprocess
import tempfile
from pathlib import Path

import chardet
import pdfplumber
from docx import Document
from .pptx_reader import read_pptx


def read_pdf(file_path):
    """
    读取PDF文件内容

    Args:
        file_path: PDF文件路径

    Returns:
        str: 提取的文本内容
    """
    text = ""

    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"

        print(f"[PDF读取] 成功读取: {Path(file_path).name}, 字数: {len(text)}")

    except Exception as e:
        print(f"[PDF读取] 读取失败 {file_path}: {e}")
        raise

    return text.strip()


def read_docx(file_path):
    """
    读取DOCX文件内容

    Args:
        file_path: DOCX文件路径

    Returns:
        str: 提取的文本内容
    """
    text = ""

    try:
        doc = Document(file_path)

        # 读取所有段落
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"

        # 读取表格中的内容
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text += cell.text + " "
            text += "\n"

        print(f"[DOCX读取] 成功读取: {Path(file_path).name}, 字数: {len(text)}")

    except Exception as e:
        print(f"[DOCX读取] 读取失败 {file_path}: {e}")
        raise

    return text.strip()


def _read_doc_via_soffice(doc_path: Path) -> str:
    """使用 LibreOffice 将 .doc 转换为 .docx 再读取。"""
    soffice_bin = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice_bin:
        return ""

    with tempfile.TemporaryDirectory(prefix="doc_to_docx_") as tmpdir:
        try:
            result = subprocess.run(
                [soffice_bin, "--headless", "--convert-to", "docx", str(doc_path), "--outdir", tmpdir],
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                timeout=60,
            )
            converted = Path(tmpdir) / f"{doc_path.stem}.docx"
            if result.returncode == 0 and converted.exists():
                return read_docx(converted)
            print(f"[DOC读取] LibreOffice 转换失败: {result.stdout or result.stderr}")
        except Exception as exc:
            print(f"[DOC读取] LibreOffice 转换异常: {exc}")

    return ""


def _read_doc_via_win32(doc_path: Path) -> str:
    """在 Windows 上通过 Word COM 读取 .doc（需要安装 Word + pywin32）。"""
    if os.name != "nt":
        return ""

    try:
        import win32com.client  # type: ignore
    except Exception as exc:
        print(f"[DOC读取] win32com 不可用: {exc}")
        return ""

    word = None
    doc = None
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(doc_path), ReadOnly=True)
        text = doc.Content.Text if doc else ""
        return text
    except Exception as exc:
        print(f"[DOC读取] win32com 读取失败: {exc}")
        return ""
    finally:
        try:
            if doc is not None:
                doc.Close(False)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass


def _read_doc_via_antiword(doc_path: Path) -> str:
    """使用 antiword CLI 读取 .doc（若已安装）。"""
    antiword_bin = shutil.which("antiword")
    if not antiword_bin:
        return ""

    try:
        result = subprocess.run(
            [antiword_bin, str(doc_path)],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=30,
        )
        text = (result.stdout or "").strip()
        if text:
            return text
        if result.stderr:
            print(f"[DOC读取] antiword 警告: {result.stderr.strip()}")
    except Exception as exc:
        print(f"[DOC读取] antiword 读取失败: {exc}")

    return ""


def read_doc(file_path):
    """
    读取DOC文件内容（老版Word）
    优先尝试直接解析，其次依次尝试 LibreOffice、Word COM、antiword。
    """
    doc_path = Path(file_path)
    last_error = None

    # 有些 .doc 实际是 OOXML 包装，先试一次 docx 解析
    try:
        return read_docx(doc_path)
    except Exception as exc:
        last_error = exc
        print(f"[DOC读取] 按 DOCX 解析失败，尝试转换: {exc}")

    for reader in (_read_doc_via_soffice, _read_doc_via_win32, _read_doc_via_antiword):
        try:
            text = reader(doc_path)
            if text and text.strip():
                print(f"[DOC读取] 使用 {reader.__name__} 成功读取: {doc_path.name}, 字数: {len(text.strip())}")
                return text.strip()
        except Exception as reader_exc:
            print(f"[DOC读取] {reader.__name__} 异常: {reader_exc}")

    raise RuntimeError(
        f"无法读取DOC文件: {doc_path}. "
        "请安装 LibreOffice (soffice) 或 antiword，或在Windows安装 Word+pywin32。 "
        f"原始解析错误: {last_error}"
    )


def read_file_content(file_path):
    """
    根据文件类型自动读取内容

    Args:
        file_path: 文件路径

    Returns:
        str: 提取的文本内容
    """
    file_path_obj = Path(file_path)
    file_ext = file_path_obj.suffix.lower()

    text_exts = {'.txt', '.md', '.html', '.htm', '.rtf', '.json', '.csv', '.py', '.js', '.ts', '.java', '.c', '.cpp', '.css'}

    if file_ext == '.pdf':
        return read_pdf(file_path)
    elif file_ext == '.docx':
        return read_docx(file_path)
    elif file_ext == '.doc':
        return read_doc(file_path)
    elif file_ext == '.pptx':
        pptx_payload = read_pptx(file_path)
        return pptx_payload.get("text", "")
    elif file_ext in text_exts:
        return read_text_file(file_path)
    else:
        raise ValueError(f"不支持的文件格式: {file_ext}")


def read_text_file(file_path, default_encoding="utf-8"):
    """
    读取纯文本/代码文件，自动尝试编码探测。
    """
    file_path_obj = Path(file_path)
    data = file_path_obj.read_bytes()
    if not data:
        return ""

    try:
        detected = chardet.detect(data)
        enc = detected.get("encoding") or default_encoding
        text = data.decode(enc, errors="ignore")
    except Exception:
        text = data.decode(default_encoding, errors="ignore")

    print(f"[文本读取] {file_path_obj.name}, 编码检测: {detected if 'detected' in locals() else 'fallback'}, 字数: {len(text)}")
    return text


def test_file_reader():
    """测试文件读取功能"""
    print("文件读取器测试")
    print("注意：此测试需要真实的PDF/DOCX文件")


if __name__ == "__main__":
    test_file_reader()
