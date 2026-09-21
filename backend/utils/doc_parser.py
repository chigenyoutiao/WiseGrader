import re
from collections import Counter
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Tuple

import mammoth
from docx import Document
from docx.shared import Length

import config


@dataclass
class FormatCheck:
    rule: str
    passed: bool
    weight: float
    message: str
    value: Optional[Any] = None


def _safe_length_pt(value: Optional[Length]) -> Optional[float]:
    if value is None:
        return None
    try:
        return float(value.pt)
    except Exception:
        return None

def _get_dominant_value(values: List[Any]) -> str:
    """获取列表中出现频率最高的值"""
    if not values:
        return "无"
    counter = Counter(values)
    most_common = counter.most_common(1)[0]
    val = most_common[0]
    # 格式化输出
    if isinstance(val, float):
        return f"{val:.1f} pt"
    return str(val)

def _detect_heading_levels(paragraphs: List[Any]) -> List[int]:
    levels = set()
    for paragraph in paragraphs:
        style = getattr(paragraph, "style", None)
        style_name = (style.name if style else "") or ""
        lower_name = style_name.lower()
        match = re.search(r"(heading|标题)\s*(\d)", lower_name)
        if match:
            try:
                levels.add(int(match.group(2)))
                continue
            except ValueError:
                pass
        text = paragraph.text.strip()
        fallback_match = re.search(r"(第?\s*\d)\s*[章|节]", text)
        if fallback_match:
            levels.add(1)
    return sorted(levels)


def _collect_paragraph_stats(paragraphs: List[Any]):
    """收集段落的物理属性统计"""
    indents = []
    spacings = []
    
    for p in paragraphs:
        if not p.text.strip():
            continue
            
        # 1. 缩进 (首行)
        indent = p.paragraph_format.first_line_indent
        if indent is None:
            indents.append("未设置")
        else:
            pt = _safe_length_pt(indent)
            if pt is not None:
                # 简单估算：2字符大约是 21-24pt (取决于字号，这里泛化处理)
                if 18 <= pt <= 30: indents.append("约2字符")
                elif pt == 0: indents.append("0字符")
                else: indents.append(f"{pt:.1f}pt")
            else:
                indents.append("未设置")

        # 2. 行距
        spacing = p.paragraph_format.line_spacing
        if spacing is None:
            spacings.append("默认(单倍)")
        elif isinstance(spacing, float):
            spacings.append(f"{spacing}倍")
        else:
            pt = _safe_length_pt(spacing)
            if pt: spacings.append(f"固定值{pt:.1f}pt")
    
    return indents, spacings


def _collect_font_sizes(paragraphs: List[Any]) -> List[float]:
    points = []
    for paragraph in paragraphs:
        # 检查 Run 级别
        for run in paragraph.runs:
            size = run.font.size
            if size:
                pt = _safe_length_pt(size)
                if pt: points.append(pt)
        # 检查 Style 级别
        if not paragraph.runs:
            style_font = getattr(paragraph.style, "font", None)
            style_size = getattr(style_font, "size", None) if style_font else None
            pt = _safe_length_pt(style_size)
            if pt: points.append(pt)
    return points


def _count_captions(paragraphs: List[Any]) -> int:
    caption_pattern = re.compile(r"^(图|Figure|Fig\.?|表)\s*\d+", re.IGNORECASE)
    return sum(1 for p in paragraphs if caption_pattern.search(p.text.strip()))


def _tables_with_header(tables: List[Any]) -> int:
    count = 0
    for table in tables:
        if not table.rows: continue
        header_row = table.rows[0]
        cell_texts = [cell.text.strip() for cell in header_row.cells]
        if any(cell_texts): count += 1
    return count


def _check_header_footer(document: Document) -> Tuple[bool, str]:
    has_content = False
    details = []
    for section in document.sections:
        h_text = " ".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
        f_text = " ".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip())
        if h_text: details.append(f"页眉: {h_text[:10]}...")
        if f_text: details.append(f"页脚: {f_text[:10]}...")
        if h_text or f_text: has_content = True
    
    return has_content, "; ".join(details) if details else "无"


def convert_docx_to_markdown(file_path: str) -> str:
    try:
        with open(file_path, "rb") as docx_file:
            result = mammoth.convert_to_markdown(docx_file)
            markdown = (result.value or "").strip()
            return markdown
    except Exception as exc:
        print(f"[DOCX] Markdown conversion failed: {exc}")
        return ""


def evaluate_docx_formatting(file_path: str) -> Dict[str, Any]:
    """
    分析 DOCX 格式，返回客观数据供 AI 评判
    """
    document = Document(file_path)
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    
    # 1. 基础统计
    paragraph_count = len(paragraphs)
    figure_count = len(document.inline_shapes)
    table_count = len(document.tables)
    
    # 2. 深度分析
    heading_levels = _detect_heading_levels(paragraphs)
    indents, spacings = _collect_paragraph_stats(paragraphs)
    font_sizes = _collect_font_sizes(paragraphs)
    has_hf, hf_detail = _check_header_footer(document)
    
    # 3. 计算“主流”格式 (Dominant Style)
    dom_indent = _get_dominant_value(indents)
    dom_spacing = _get_dominant_value(spacings)
    dom_font = _get_dominant_value(font_sizes)

    # 4. 构建数据概览 (Metrics)
    metrics = {
        "paragraph_count": paragraph_count,
        "heading_levels": heading_levels if heading_levels else "未检测到标题样式",
        "dominant_indent": dom_indent,
        "dominant_spacing": dom_spacing,
        "dominant_font_size": dom_font,
        "header_footer_status": "存在" if has_hf else "未设置",
        "header_footer_detail": hf_detail,
        "figure_count": figure_count,
        "table_count": table_count,
        "caption_count": _count_captions(paragraphs)
    }

    # 5. 提取文本内容
    markdown = convert_docx_to_markdown(file_path)
    
    return {
        "score": 0, # 不再由代码打分，交由 AI
        "checks": [], # 不再生成硬编码检查项
        "metrics": metrics,
        "markdown": markdown,
        "confidence": 1.0,
    }